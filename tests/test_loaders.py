from __future__ import annotations

import io

import pytest
from docx import Document
from pypdf import PdfWriter

from docpilot.loaders import DocumentLoader, UnsupportedDocumentError


def test_txt_loader_supports_chinese() -> None:
    docs = DocumentLoader().load_bytes("向量数据库".encode(), "notes.txt")
    assert docs[0].text == "向量数据库"


def test_markdown_loader_treats_markdown_as_text() -> None:
    docs = DocumentLoader().load_bytes("# 标题\n\n正文".encode(), "guide.md")
    assert docs[0].metadata["file_type"] == "md"
    assert "正文" in docs[0].text


def test_csv_loader_keeps_header_names() -> None:
    docs = DocumentLoader().load_bytes("课程,教师\n人工智能,小王\n".encode(), "courses.csv")
    assert len(docs) == 1
    assert "课程: 人工智能" in docs[0].text
    assert docs[0].metadata["row"] == 2


def test_docx_loader_reads_paragraphs_and_tables() -> None:
    buffer = io.BytesIO()
    document = Document()
    document.add_paragraph("项目说明")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "技术"
    table.cell(0, 1).text = "Milvus"
    document.save(buffer)
    docs = DocumentLoader().load_bytes(buffer.getvalue(), "project.docx")
    assert "项目说明" in docs[0].text
    assert "技术 | Milvus" in docs[0].text


def test_pdf_loader_accepts_a_valid_pdf() -> None:
    buffer = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.write(buffer)
    assert DocumentLoader().load_bytes(buffer.getvalue(), "blank.pdf") == []


def test_unsupported_file_type() -> None:
    with pytest.raises(UnsupportedDocumentError):
        DocumentLoader().load_bytes(b"data", "archive.zip")
