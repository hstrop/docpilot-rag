from __future__ import annotations

import csv
import io
from pathlib import Path

from docpilot.domain import LoadedDocument

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".docx", ".csv"}


class UnsupportedDocumentError(ValueError):
    pass


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("unable to decode text file as UTF-8 or GB18030")


class DocumentLoader:
    def load_path(self, path: str | Path) -> list[LoadedDocument]:
        source_path = Path(path).expanduser().resolve()
        if not source_path.is_file():
            raise FileNotFoundError(str(source_path))
        return self.load_bytes(source_path.read_bytes(), source_path.name)

    def load_bytes(self, content: bytes, filename: str) -> list[LoadedDocument]:
        suffix = Path(filename).suffix.lower()
        if suffix not in SUPPORTED_EXTENSIONS:
            supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
            raise UnsupportedDocumentError(
                f"unsupported file type {suffix!r}; expected {supported}"
            )
        if suffix == ".txt":
            return [LoadedDocument(_decode_text(content), filename, {"file_type": "txt"})]
        if suffix == ".csv":
            return self._load_csv(content, filename)
        if suffix == ".docx":
            return self._load_docx(content, filename)
        return self._load_pdf(content, filename)

    @staticmethod
    def _load_csv(content: bytes, filename: str) -> list[LoadedDocument]:
        text = _decode_text(content)
        reader = csv.DictReader(io.StringIO(text))
        if not reader.fieldnames:
            return [LoadedDocument(text, filename, {"file_type": "csv"})]
        documents = []
        for row_number, row in enumerate(reader, start=2):
            line = " | ".join(f"{key}: {value or ''}" for key, value in row.items())
            documents.append(
                LoadedDocument(line, filename, {"file_type": "csv", "row": row_number})
            )
        return documents

    @staticmethod
    def _load_docx(content: bytes, filename: str) -> list[LoadedDocument]:
        try:
            from docx import Document
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("python-docx is required to read DOCX files") from exc
        document = Document(io.BytesIO(content))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells))
        return [LoadedDocument("\n".join(paragraphs), filename, {"file_type": "docx"})]

    @staticmethod
    def _load_pdf(content: bytes, filename: str) -> list[LoadedDocument]:
        try:
            from pypdf import PdfReader
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("pypdf is required to read PDF files") from exc
        reader = PdfReader(io.BytesIO(content))
        documents = []
        for page_number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                documents.append(
                    LoadedDocument(
                        text,
                        filename,
                        {"file_type": "pdf", "page": page_number},
                    )
                )
        return documents
